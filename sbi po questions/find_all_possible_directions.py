import os
import re
import docx

sbi_docx_dir = r"C:\Users\Administrator\Downloads\sbi po prelims"

def iter_block_items(parent):
    if isinstance(parent, docx.document.Document):
        parent_elm = parent.element.body
    elif isinstance(parent, docx.table._Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent type")
    for child in parent_elm.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            yield docx.text.paragraph.Paragraph(child, parent)
        elif tag == 'tbl':
            yield docx.table.Table(child, parent)

def get_rel_map(doc):
    rel_map = {}
    for rel_id, rel in doc.part.rels.items():
        if "image" in rel.target_ref:
            rel_map[rel_id] = rel.target_ref
    return rel_map

def get_paragraph_text_with_images(p, rel_map):
    text = p.text.strip()
    blips = p._p.xpath('.//a:blip')
    p_img_placeholders = []
    for blip in blips:
        from docx.oxml.ns import qn
        embed_id = blip.get(qn('r:embed'))
        if embed_id in rel_map:
            img_name = os.path.basename(rel_map[embed_id])
            p_img_placeholders.append(f"[IMAGE:{img_name}]")
    if p_img_placeholders:
        if text:
            text += " " + " ".join(p_img_placeholders)
        else:
            text = " ".join(p_img_placeholders)
    return text

def is_layout_table(table):
    if len(table.rows) == 1:
        return True
    for row in table.rows:
        for cell in row.cells:
            txt = cell.text.strip()
            if re.search(r'\bQ\d+\b', txt) or re.search(r'\bS\d+\s*\.\s*Ans', txt, re.IGNORECASE):
                return True
    return False

def format_data_table_as_markdown(table, rel_map):
    rows_data = []
    for row in table.rows:
        row_cells = []
        for cell in row.cells:
            cell_paras = []
            for p in cell.paragraphs:
                p_txt = get_paragraph_text_with_images(p, rel_map)
                if p_txt:
                    cell_paras.append(p_txt)
            row_cells.append(" ".join(cell_paras).strip())
        rows_data.append(row_cells)
    if not rows_data or len(rows_data[0]) == 0:
        return ""
    md_lines = []
    header = rows_data[0]
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in rows_data[1:]:
        if len(row) < len(header):
            row.extend([""] * (len(header) - len(row)))
        elif len(row) > len(header):
            row = row[:len(header)]
        md_lines.append("| " + " | ".join(row) + " |")
    return "\n" + "\n".join(md_lines) + "\n"

def extract_content_blocks(parent, rel_map):
    blocks = []
    for item in iter_block_items(parent):
        if isinstance(item, docx.text.paragraph.Paragraph):
            txt = get_paragraph_text_with_images(item, rel_map)
            if txt:
                blocks.append(txt)
        elif isinstance(item, docx.table.Table):
            if is_layout_table(item):
                for col_idx in range(len(item.columns)):
                    for row_idx in range(len(item.rows)):
                        cell = item.cell(row_idx, col_idx)
                        blocks.extend(extract_content_blocks(cell, rel_map))
            else:
                md_table = format_data_table_as_markdown(item, rel_map)
                if md_table:
                    blocks.append(md_table)
    return blocks

def find_directions():
    files = sorted([f for f in os.listdir(sbi_docx_dir) if f.lower().endswith('.docx')])
    for idx, file_name in enumerate(files):
        test_num = idx + 1
        file_path = os.path.join(sbi_docx_dir, file_name)
        doc = docx.Document(file_path)
        rel_map = get_rel_map(doc)
        blocks = extract_content_blocks(doc, rel_map)
        
        # Split into Questions and Solutions
        sol_start_idx = None
        for i, b in enumerate(blocks):
            txt = b.strip()
            if re.match(r'^S1\s*\.\s*Ans', txt, re.IGNORECASE) or txt.lower() in ["solutions", "answers", "detailed solutions"]:
                sol_start_idx = i
                break
        if sol_start_idx is None:
            for i, b in enumerate(blocks):
                if re.search(r'\bS1\s*\.\s*Ans', b, re.IGNORECASE):
                    sol_start_idx = i
                    break
        if sol_start_idx is None:
            sol_start_idx = len(blocks)
            
        questions_text = "\n\n".join(blocks[:sol_start_idx])
        
        # Search for any variation of Direction / Directions (Q. X-Y) or (X-Y)
        pattern = r'\bDirections?\b[^\n]{0,100}'
        matches = re.findall(pattern, questions_text, re.IGNORECASE)
        print(f"\nFile {test_num}: {file_name}")
        for m in matches:
            # check if it has a range of numbers like 1-5 or similar
            if re.search(r'\d+', m):
                print(f"  Match: {m.strip()}")

if __name__ == "__main__":
    find_directions()
